"""Parser node: extract the machine-readable text layer for the file converter.py already
turned into page images. Docling first (cheap, CPU); falls back to the byproduct text
converter.py captured while rendering pages, or a pure-python DOCX reader, when Docling is
unavailable or yields nothing.
"""
import json
from functools import lru_cache
from pathlib import Path

from agent.logger import get_logger, log_node
from agent.state import HRState, RDEMError
from config import settings as S
from utils.text import alnum_ratio, presentation_form_ratio

_TXT = {".txt", ".md", ".json"}


def text_is_usable(text: str) -> bool:
    """Router check from the plan: near-empty, noisy, or garbled-Arabic text layers escalate to vision."""
    return (
        len(text.strip()) >= S.MIN_TEXT_CHARS
        and alnum_ratio(text) >= S.MIN_ALNUM_RATIO
        and presentation_form_ratio(text) <= S.MAX_PRESENTATION_FORM_RATIO
    )


@lru_cache(maxsize=1)
def _converter():
    """One converter per process (building it loads models). Mode comes from DOCLING_MODE.

    Why not the library default: it runs Docling's OCR on scanned pages. That OCR text can look
    'usable' to text_is_usable() (long enough, alphanumeric) while being wrong, especially for
    Arabic, and then the VLM would never be asked. With OCR off, a scan yields no text and the
    router escalates to the VLM, which is the design.
    """
    from docling.datamodel.base_models import InputFormat  # lazy, heavy
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    mode = S.DOCLING_MODE
    if mode == "full":
        return DocumentConverter()
    opts = PdfPipelineOptions(do_ocr=False, do_table_structure=(mode == "tables"))
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})


def _docling_text(path: Path) -> str:
    try:
        return _converter().convert(str(path)).document.export_to_markdown()
    except Exception as e:
        get_logger().debug(f"    docling unavailable/failed ({type(e).__name__}); using fallback")
        return ""


def _docx_text_fallback(path: Path) -> str:
    """Pure-python text for .docx (paragraphs + table cells) when Docling is unavailable."""
    try:
        from docx import Document  # python-docx (already a Docling dependency)

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n".join(parts)
    except Exception as e:
        get_logger().debug(f"    python-docx fallback failed ({type(e).__name__})")
        return ""


@log_node
def parser(state: HRState) -> dict:
    path = Path(state.cv_path)
    ext = path.suffix.lower()
    text = ""

    try:
        if ext in _TXT:
            raw = path.read_text(encoding="utf-8")
            text = json.dumps(json.loads(raw), ensure_ascii=False) if ext == ".json" else raw
        elif ext == ".pdf":
            text = _docling_text(path) or state.raw_text_layer
        elif ext in {".docx", ".doc"}:
            text = _docling_text(path) or (_docx_text_fallback(path) if ext == ".docx" else "")
            text = text or state.raw_text_layer
        # image files have no separate text layer; text stays ""
    except Exception as e:
        rdem = RDEMError(node="parser", error_type="parse_failed", message=f"{type(e).__name__}: {e}",
                         suggestion="Check the file is a valid, non-encrypted CV.")
        return {"status": "failed", "global_error_log": [rdem], "_status": "error"}

    if not state.page_paths and not text.strip():
        rdem = RDEMError(node="parser", error_type="empty_document",
                         message="No text or page images could be produced.")
        return {"status": "failed", "global_error_log": [rdem], "_status": "error"}

    ok = text_is_usable(text)
    get_logger().info(f"  [parser] id={state.candidate_id} text_chars={len(text)} text_ok={ok}")
    return {"text_layer": text, "text_ok": ok}
